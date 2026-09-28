"""Build docs/outbox/Thesis_Muller_v3.docx from thesis_full.md.

Applies the programme's formatting rules: Times New Roman 12 pt, 2 cm margins,
1.5 line spacing, page numbers. Each chapter starts on a new page; the Contents
list becomes a Word TOC field (Word asks to update it on open).
Run build.py first so thesis_full.md is current.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
SRC = HERE / "thesis_full.md"
OUT = HERE.parent / "outbox" / "Thesis_Muller_v3.docx"
FONT = "Times New Roman"


def field(par, instr):
    """Insert a Word field (PAGE, TOC ...) into a paragraph."""
    run = par.add_run()
    for tag, text in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def add_inline(par, text):
    """Markdown **bold**, *italic*, `code` -> runs."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("*"):
            par.add_run(tok[1:-1]).italic = True
        elif tok.startswith("`"):
            par.add_run(tok[1:-1])
        else:
            par.add_run(tok)


def set_font(style):
    style.font.name = FONT
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


def setup(doc):
    st = doc.styles["Normal"]
    set_font(st)
    st.font.size = Pt(12)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 18), ("Heading 1", 16), ("Heading 2", 13)):
        h = doc.styles[name]
        set_font(h)
        h.font.size = Pt(size)
        h.font.color.rgb = None
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2)
        p = s.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field(p, "PAGE")
    # ask Word to refresh the TOC and page fields when the file opens
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    doc.settings.element.append(upd)


def table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-+:?", c) for c in r)]
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Table Grid"
    for i, r in enumerate(cells):
        for j, c in enumerate(r):
            par = t.cell(i, j).paragraphs[0]
            par.paragraph_format.line_spacing = 1.0
            add_inline(par, c)
            for run in par.runs:
                run.font.size = Pt(10)
                run.bold = run.bold or i == 0


def main():
    text = re.sub(r"<!--.*?-->", "", SRC.read_text(encoding="utf-8"), flags=re.S)
    doc = Document()
    setup(doc)
    first_h1, skip_contents = True, False
    for b in re.split(r"\n\s*\n", text):
        lines = b.strip().split("\n")
        head = lines[0]
        if not head.strip() or head.strip() == "---":
            continue
        if skip_contents:  # drop the hand-written Contents list; the TOC replaces it
            if not head.startswith("#"):
                continue
            skip_contents = False
        if head.startswith("# "):
            if first_h1:
                p = doc.add_paragraph(style="Title")
                add_inline(p, head[2:])
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first_h1 = False
            else:
                doc.add_page_break()
                doc.add_heading(head[2:], level=1)
            continue
        if head.startswith("## "):
            if head[3:] == "Contents":
                doc.add_page_break()
                doc.add_heading("Contents", level=1)
                field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
                skip_contents = True
            else:
                doc.add_heading(head[3:], level=2)
            continue
        if head.lstrip().startswith("|"):
            table(doc, lines)
            continue
        m = re.match(r"!\[(.*?)\]\((.*?)\)", head)
        if m:
            doc.add_picture(str((HERE / m.group(2)).resolve()), width=Cm(15))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(doc.add_paragraph(), "*" + m.group(1) + "*")
            continue
        if re.match(r"\s*(- |\d+\. )", head):
            items = []
            for l in lines:
                if re.match(r"\s*(- |\d+\. )", l):
                    items.append(re.sub(r"^\s*(- |\d+\. )", "", l))
                else:
                    items[-1] += " " + l.strip()
            style = "List Bullet" if head.lstrip().startswith("- ") else "List Number"
            for it in items:
                add_inline(doc.add_paragraph(style=style), it)
            continue
        p = doc.add_paragraph()
        if head.startswith(("**Student", "*Data Quality e")):
            # title-page block: one field per line, centred
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for i, l in enumerate(lines):
                if i:
                    p.add_run().add_break(WD_BREAK.LINE)
                add_inline(p, l.strip())
            if head.startswith("**Student"):
                doc.add_page_break()
            continue
        add_inline(p, " ".join(l.strip() for l in lines))
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
